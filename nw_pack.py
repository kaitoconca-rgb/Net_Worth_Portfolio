"""Accountant pack: one Excel workbook per Australian financial year.

Sep 2026. Pulls together what the tax agent asks for each year:
  * every rental property in Zarpia (live or the saved copy): one sheet each,
    plus portfolio totals on the overview
  * foreign investment income: interest, dividends, bond coupons (income ledger)
  * realised gains and losses: sales and maturities (Cost base & gains)
  * the Reserve Bank exchange rates used
All A$ figures use Reserve Bank of Australia rates: the payment/sale date
rate for single items, the ATO financial-year average (in each property's own
currency) for rental properties (income and costs spread through the year).
"""
import io
from datetime import date, datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import nw_zarpia
from nw_income import TYPE_LABELS, load_income
from nw_lots import realised_gains

F = "Arial"
TITLE = Font(name=F, size=14, bold=True)
BOLD = Font(name=F, bold=True)
NORM = Font(name=F)
NOTE = Font(name=F, italic=True, color="555555")
HEAD = Font(name=F, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2F5D50")
SUB_FILL = PatternFill("solid", fgColor="E8EFEC")
WARN = Font(name=F, color="B00000")
MONEY = '#,##0.00;-#,##0.00;"-"'
RATE = "0.0000"
DATE = "dd/mm/yyyy"


def fy_bounds(fy_year):
    return date(fy_year - 1, 7, 1), date(fy_year, 6, 30)


def fy_name(fy_year):
    return f"FY{str(fy_year)[-2:]}"


def current_fy():
    t = date.today()
    return t.year + 1 if t.month >= 7 else t.year


# ─────────────────────────── gathering the data ─────────────────────────────

def income_for_fy(conn, fy_year, aud_rate_on):
    df = load_income(conn)
    start, end = fy_bounds(fy_year)
    d = df.copy()
    d["div_date"] = pd.to_datetime(d["div_date"]).dt.date
    d = d[(d["div_date"] >= start) & (d["div_date"] <= end)].copy()
    if d.empty:
        return d
    d["Type"] = d["income_type"].map(TYPE_LABELS).fillna("Dividend")
    d["Country"] = d["country"].fillna("?")
    d["Gross"] = d["gross_amount"].fillna(d["amount"])
    d["Tax withheld"] = d["tax_withheld"].fillna(0.0)
    d["rate_missing"] = d["fx_rate_to_aud"].isna()
    d["AUD rate"] = [
        1.0 if str(c).upper() == "AUD" else (r if pd.notna(r) else aud_rate_on(c, dt))
        for c, r, dt in zip(d["currency"], d["fx_rate_to_aud"], d["div_date"])]
    d["AUD rate"] = pd.to_numeric(d["AUD rate"], errors="coerce")
    d["Gross A$"] = d["Gross"] * d["AUD rate"]
    d["Tax withheld A$"] = d["Tax withheld"] * d["AUD rate"]
    d["Net A$"] = d["amount"] * d["AUD rate"]
    d["no_gross"] = d["gross_amount"].isna()
    return d.sort_values("div_date")


def gains_for_fy(conn, fy_year, aud_rate_on):
    g = realised_gains(conn, aud_rate_on)
    if g is None or g.empty:
        return pd.DataFrame()
    return g[g["FY"] == fy_name(fy_year)].copy()


def property_for_fy(pg, fy_year):
    """([summary per property with activity in the FY], source text) or ([], reason)."""
    try:
        if nw_zarpia.feed_configured():
            feed, src = nw_zarpia.load_feed(), f"Zarpia (live, read {datetime.now():%d %b %Y})"
        else:
            feed = nw_zarpia.load_snapshot(pg)
            if feed is None:
                return [], "Zarpia not connected"
            src = f"Zarpia (copy of {feed['copied_at']:%d %b %Y})"
    except Exception as e:
        return [], f"Couldn't read Zarpia: {e}"
    props = nw_zarpia.portfolio_fy_summaries(feed, fy_year)
    if not props:
        return [], f"No Zarpia rent or expenses for {fy_name(fy_year)}"
    return props, src


def _sheet_title(name, used):
    """Excel sheet name: <= 31 chars, no []:*?/\ , unique in the workbook."""
    base = "".join(ch for ch in f"Rental - {name}" if ch not in '[]:*?/\\')[:31]
    t, i = base, 2
    while t in used:
        t = f"{base[:28]} {i}"
        i += 1
    used.add(t)
    return t


# ─────────────────────────── writing the workbook ───────────────────────────

def _w(ws, r, c, v, fmt=None, font=None, fill=None, wrap=False):
    if isinstance(v, pd.Timestamp):
        v = v.date()
    try:
        if v is not None and not isinstance(v, (str, date)) and pd.isna(v):
            v = None
    except (TypeError, ValueError):
        pass
    cell = ws.cell(r, c, v)
    cell.font = font or NORM
    if isinstance(v, date):
        cell.number_format = DATE
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    cell.alignment = Alignment(vertical="top", wrap_text=wrap)
    return cell


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _table(ws, r, df, cols, fmts=None, total_cols=(), total_label="Total"):
    """Header + rows (+ a SUM total row). Returns the next free row."""
    fmts = fmts or {}
    for j, c in enumerate(cols, 1):
        _w(ws, r, j, c, font=HEAD, fill=HEAD_FILL, wrap=True)
    first = r + 1
    for i, (_, row) in enumerate(df.iterrows()):
        for j, c in enumerate(cols, 1):
            _w(ws, first + i, j, row.get(c), fmts.get(c))
    last = first + len(df) - 1
    r = last + 1
    if total_cols and len(df):
        _w(ws, r, 1, total_label, font=BOLD, fill=SUB_FILL)
        for j, c in enumerate(cols, 1):
            if j > 1:
                ws.cell(r, j).fill = SUB_FILL
            if c in total_cols:
                col = get_column_letter(j)
                _w(ws, r, j, f"=SUM({col}{first}:{col}{last})", fmts.get(c, MONEY), BOLD, SUB_FILL)
        r += 1
    return r


def _property_sheet(wb, prop, fy_year, prop_src, rate, points):
    fy = fy_name(fy_year)
    ccy = prop["currency"]
    show_aud = bool(rate) and ccy != "AUD"
    eur_avg = rate if show_aud else None
    ws = wb.create_sheet(prop["sheet"])
    ws["A1"] = f"Rental property - {prop['name']}, {prop['country_name']} - {fy}"
    ws["A1"].font = TITLE
    ws["A2"] = (f"Source: {prop_src}. {ccy} amounts; A$ at the ATO's {fy} average rate "
                f"({1 / eur_avg:.4f} {ccy} per A$, i.e. {eur_avg:.4f} A$ per {ccy})." if eur_avg
                else f"Source: {prop_src}. {ccy} amounts.")
    ws["A2"].font = NOTE
    r = 4
    rows = [("Weeks available for rent", prop["weeks_available"], None,
             f"{prop['owner_nights']} owner-use nights excluded"),
            ("Weeks rented", prop["weeks_rented"], None, f"{prop['rented_nights']} booked nights ÷ 7"),
            ("Gross rent", prop["rent"], "A$", "Bookings, for nights stayed in the year"),
            ("Deductible expenses", prop["expenses_total"], "A$", "By category below")]
    for j, h in enumerate(["Item", f"Value ({ccy} / weeks)", "A$", "Note"], 1):
        _w(ws, r, j, h, font=HEAD, fill=HEAD_FILL)
    r += 1
    for label, v, aud, note in rows:
        _w(ws, r, 1, label)
        _w(ws, r, 2, v, "0.0" if aud is None else MONEY)
        if aud and eur_avg:
            _w(ws, r, 3, f"=B{r}*{round(eur_avg, 6)}", MONEY)
        _w(ws, r, 4, note)
        r += 1
    _w(ws, r, 1, "Net rent before depreciation", font=BOLD)
    _w(ws, r, 2, f"=B{r - 2}-B{r - 1}", MONEY, BOLD)
    if eur_avg:
        _w(ws, r, 3, f"=C{r - 2}-C{r - 1}", MONEY, BOLD)
    r += 2

    _w(ws, r, 1, "Expenses by category", font=BOLD)
    r += 1
    cat = prop["expenses_by_cat"].reset_index()
    cat.columns = ["Category", ccy]
    cat["A$"] = cat[ccy] * eur_avg if eur_avg else None
    r = _table(ws, r, cat, ["Category", ccy, "A$"], {ccy: MONEY, "A$": MONEY}, (ccy, "A$")) + 1

    if prop["spanish_tax"]:
        _w(ws, r, 1, "Spanish non-resident tax on the rent (24% of gross, Modelo 210)", font=BOLD)
        r += 1
        sp = pd.DataFrame([
            {"Period": f"Jul-Dec {fy_year - 1}", "Rent (EUR)": prop["rent_h2"], "Tax (EUR)": prop["spanish_tax_h2"],
             "Tax A$": prop["spanish_tax_h2"] * eur_avg if eur_avg else None,
             "Filed with": f"{fy_year - 1} Spanish returns (by Jan {fy_year})"},
            {"Period": f"Jan-Jun {fy_year}", "Rent (EUR)": prop["rent_h1"], "Tax (EUR)": prop["spanish_tax_h1"],
             "Tax A$": prop["spanish_tax_h1"] * eur_avg if eur_avg else None,
             "Filed with": f"{fy_year} Spanish returns (by Jan {fy_year + 1})"}])
        r = _table(ws, r, sp, ["Period", "Rent (EUR)", "Tax (EUR)", "Tax A$", "Filed with"],
                   {"Rent (EUR)": MONEY, "Tax (EUR)": MONEY, "Tax A$": MONEY}, ("Rent (EUR)", "Tax (EUR)", "Tax A$"))
        _w(ws, r, 1, "Calculated on gross rent. Tax on deemed income for owner-use days and the amounts actually "
                     "paid come from the Modelo 210 filings.", font=NOTE)
        r += 2
    else:
        _w(ws, r, 1, f"Rent by calendar half ({prop['country_name']} taxes by calendar year)", font=BOLD)
        r += 1
        hh = pd.DataFrame([
            {"Period": f"Jul-Dec {fy_year - 1}", f"Rent ({ccy})": prop["rent_h2"],
             "Rent A$": prop["rent_h2"] * eur_avg if eur_avg else None},
            {"Period": f"Jan-Jun {fy_year}", f"Rent ({ccy})": prop["rent_h1"],
             "Rent A$": prop["rent_h1"] * eur_avg if eur_avg else None}])
        r = _table(ws, r, hh, ["Period", f"Rent ({ccy})", "Rent A$"], {f"Rent ({ccy})": MONEY, "Rent A$": MONEY},
                   (f"Rent ({ccy})", "Rent A$"))
        _w(ws, r, 1, f"Tax paid in {prop['country_name']}: from the local returns (see Zarpia's Tax page).",
           font=NOTE)
        r += 2

    det = prop["deductible_detail"].rename(columns={
        "expense_date": "Date", "supplier": "Supplier", "concept": "Description", "unit": "Unit",
        "amount": ccy, "pdf_filename": "Document"})
    _w(ws, r, 1, "Expense detail", font=BOLD)
    r += 1
    r = _table(ws, r, det, ["Date", "Accountant category", "Supplier", "Description", "Unit", ccy, "Document"],
               {ccy: MONEY}, (ccy,)) + 1
    cap = prop["capital_items"]
    if not cap.empty:
        _w(ws, r, 1, "Capital items (depreciate - not an immediate deduction)", font=BOLD)
        r += 1
        c2 = cap.rename(columns={"expense_date": "Date", "concept": "Description", "unit": "Unit", "amount": ccy})
        r = _table(ws, r, c2, ["Date", "Description", "Unit", ccy], {ccy: MONEY}, (ccy,)) + 1
        points.append(f"{prop['name']}: {len(cap)} capital item(s) ({ccy} {cap['amount'].sum():,.2f}) - decline in "
                      "value to be calculated.")
    exc = prop["excluded"]
    if not exc.empty:
        _w(ws, r, 1, "Left out of deductions (private or not claimable)", font=BOLD)
        r += 1
        e2 = exc.rename(columns={"expense_date": "Date", "category": "Category", "concept": "Description",
                                 "amount": ccy})
        r = _table(ws, r, e2, ["Date", "Category", "Description", ccy], {ccy: MONEY}, (ccy,)) + 1
    if len(prop["bookings"]):
        bk = prop["bookings"].rename(columns={"booking_ref": "Booking", "arrival": "Arrival",
                                              "departure": "Departure", "nights": "Nights",
                                              "nights_fy": f"Nights in {fy}", "income_total": f"Rent ({ccy})",
                                              "agent_statement": "Agent statement"})
        bk["Owner use"] = bk["rented"].map({True: "", False: "Owner use"})
        _w(ws, r, 1, "Bookings", font=BOLD)
        r += 1
        r = _table(ws, r, bk, ["Booking", "Arrival", "Departure", "Nights", f"Nights in {fy}", f"Rent ({ccy})",
                               "Agent statement", "Owner use"], {f"Rent ({ccy})": MONEY})
        r += 1
    mi = prop["manual_income"]
    if len(mi):
        m2 = mi.rename(columns={"source": "Source", "checkin": "Check-in", "checkout": "Check-out",
                                "nights": "Nights", "amount": f"Amount ({ccy})",
                                "amount in FY": f"In {fy} ({ccy})", "nights in FY": f"Nights in {fy}"})
        _w(ws, r, 1, "Other income (manual / email)", font=BOLD)
        r += 1
        r = _table(ws, r, m2, ["Source", "Check-in", "Check-out", "Nights", f"Amount ({ccy})", f"In {fy} ({ccy})",
                               f"Nights in {fy}"], {f"Amount ({ccy})": MONEY, f"In {fy} ({ccy})": MONEY},
                   (f"In {fy} ({ccy})",)) + 1
    if len(prop["undated_income"]):
        points.append(f"{prop['name']}: {len(prop['undated_income'])} income entry(ies) in Zarpia have no dates "
                      "and are not counted - add dates in Zarpia.")
    _widths(ws, [30, 24, 22, 40, 20, 14, 30, 12])
    levy = float(prop["expenses_by_cat"].get("Strata special levy", 0) or 0)
    if levy:
        points.append(f"{prop['name']}: special levies (derrama) {ccy} {levy:,.2f} included in expenses - purpose "
                      "to confirm (capital if they fund improvements).")
    if prop["spanish_tax"]:
        points.append(f"{prop['name']}: Spanish tax is paid by calendar year - the Jan-Jun part is paid the "
                      "following January; foreign income tax offset timing to confirm.")


def build_workbook(fy_year, props, prop_src, rate_of, income, gains, rates_used):
    fy = fy_name(fy_year)
    start, end = fy_bounds(fy_year)
    wb = Workbook()
    ov = wb.active
    ov.title = "Overview"
    points = []

    # ── Rental properties (one sheet each) ──
    used = set()
    for prop in props:
        prop["sheet"] = _sheet_title(prop["name"], used)
        _property_sheet(wb, prop, fy_year, prop_src, rate_of(prop["currency"]), points)

    # ── Income ──
    if income is not None and not income.empty:
        ws = wb.create_sheet("Investment income")
        ws["A1"] = f"Interest, dividends and bond coupons - {fy}"
        ws["A1"].font = TITLE
        ws["A2"] = ("Counted when paid into the account. A$ at the Reserve Bank rate on the payment date. "
                    "Australian items (country AU) are pre-filled by the ATO - shown for completeness.")
        ws["A2"].font = NOTE
        inc = income.rename(columns={"div_date": "Date", "portfolio": "Account", "payer": "Payer",
                                     "security": "Security", "amount": "Net", "currency": "Currency"})
        r = _table(ws, 4, inc, ["Date", "Type", "Account", "Payer", "Country", "Security", "Currency", "Gross",
                                "Tax withheld", "Net", "AUD rate", "Gross A$", "Tax withheld A$", "Net A$"],
                   {"Gross": MONEY, "Tax withheld": MONEY, "Net": MONEY, "AUD rate": RATE, "Gross A$": MONEY,
                    "Tax withheld A$": MONEY, "Net A$": MONEY},
                   ("Gross A$", "Tax withheld A$", "Net A$"))
        _widths(ws, [12, 12, 14, 28, 9, 26, 9, 12, 12, 12, 10, 12, 14, 12])
        ws.freeze_panes = "A5"
        if income["no_gross"].any():
            points.append(f"Income: {int(income['no_gross'].sum())} payment(s) have no gross/tax recorded - net "
                          "shown as gross; check the provider's annual tax statement.")
        if income["rate_missing"].any():
            points.append(f"Income: {int(income['rate_missing'].sum())} payment(s) had no stored rate - converted "
                          "with the Reserve Bank rate for that date.")

    # ── Gains ──
    if gains is not None and not gains.empty:
        ws = wb.create_sheet("Gains and losses")
        ws["A1"] = f"Sales and maturities - {fy}"
        ws["A1"].font = TITLE
        ws["A2"] = ("Gain = proceeds in A$ (rate on the disposal date) less cost in A$ (rate on the acquisition date). "
                    "Purchases matched first-in first-out. Inherited parcels: cost = market value at date of death.")
        ws["A2"].font = NOTE
        cols = ["Asset", "ISIN", "Source", "How acquired", "Acquired", "How disposed", "Disposed", "Quantity",
                "Currency", "Cost", "Cost FX", "Cost A$", "Proceeds", "Proceeds FX", "Proceeds A$",
                "Gain / (loss) A$", "Held > 12 months"]
        g = gains.copy()
        g["Held > 12 months"] = g["Held > 12 months"].map({True: "Yes", False: "No"}).fillna("?")
        _table(ws, 4, g, cols, {"Quantity": "#,##0.####", "Cost": MONEY, "Cost FX": RATE, "Cost A$": MONEY,
                                 "Proceeds": MONEY, "Proceeds FX": RATE, "Proceeds A$": MONEY,
                                 "Gain / (loss) A$": MONEY}, ("Cost A$", "Proceeds A$", "Gain / (loss) A$"))
        _widths(ws, [26, 15, 16, 22, 11, 16, 11, 11, 9, 12, 9, 12, 12, 10, 12, 14, 10])
        ws.freeze_panes = "A5"
        unknown = g["Cost A$"].isna().sum()
        if unknown:
            points.append(f"Gains: {int(unknown)} sale(s) have no matching purchase - cost base missing, add the "
                          "parcel in Cost base & gains.")
        if (g["How disposed"].astype(str).str.contains("atur", na=False)).any():
            points.append("Gains: bonds redeemed at maturity - confirm capital gain vs ordinary income "
                          "(traditional security / foreign-exchange rules).")
        losses = g.loc[g["Gain / (loss) A$"] < 0, "Gain / (loss) A$"].sum()
        if losses:
            points.append(f"Gains: capital losses A${-losses:,.2f} offset capital gains; any excess carries forward.")

    # ── FX rates ──
    ws = wb.create_sheet("Exchange rates")
    ws["A1"] = "Exchange rates used"
    ws["A1"].font = TITLE
    ws["A2"] = ("Reserve Bank of Australia, statistical table F11.1 (4pm AEST), shown as A$ per 1 unit of foreign "
                "currency. Weekends and holidays use the previous business day.")
    ws["A2"].font = NOTE
    rt = pd.DataFrame(rates_used, columns=["Date", "Currency", "A$ per unit", "Used for"])
    _table(ws, 4, rt, ["Date", "Currency", "A$ per unit", "Used for"], {"A$ per unit": RATE})
    _widths(ws, [14, 10, 12, 70])

    # ── Overview ──
    ov["A1"] = f"Foreign income summary - {fy} (1 Jul {fy_year - 1} - 30 Jun {fy_year})"
    ov["A1"].font = TITLE
    ov["A2"] = (f"Prepared {date.today():%d %b %Y} from Claudio's Executive Console. Australian salary, dividends and "
                "bank interest are pre-filled by the ATO and not repeated here. For discussion with the tax agent - "
                "not tax advice.")
    ov["A2"].font = NOTE
    ov["A2"].alignment = Alignment(wrap_text=True)
    ov.merge_cells("A2:F2")
    ov.row_dimensions[2].height = 32
    r = 4
    for j, h in enumerate(["Section", "Item", "A$", "Foreign tax paid A$", "Detail sheet", "Note"], 1):
        _w(ov, r, j, h, font=HEAD, fill=HEAD_FILL)
    r += 1

    def sec(title):
        nonlocal r
        _w(ov, r, 1, title, font=BOLD)
        for j in range(1, 7):
            ov.cell(r, j).fill = SUB_FILL
        r += 1

    def line(item, aud=None, tax=None, sheet="", note="", bold=False, fmt=MONEY):
        nonlocal r
        f = BOLD if bold else NORM
        _w(ov, r, 2, item, font=f)
        if aud is not None:
            _w(ov, r, 3, round(float(aud), 2), fmt, f)
        if tax is not None:
            _w(ov, r, 4, round(float(tax), 2), MONEY, f)
        _w(ov, r, 5, sheet)
        _w(ov, r, 6, note, wrap=len(note) > 60)
        r += 1

    sec("1. Rental properties")
    if props:
        tot_rent = tot_exp = tot_tax = 0.0
        missing_rate = False
        for prop in props:
            rt = rate_of(prop["currency"])
            if not rt:
                missing_rate = True
                line(f"{prop['name']} ({prop['country_name']})", sheet=prop["sheet"],
                     note=f"No ATO average rate for {prop['currency']} - see sheet ({prop['currency']} amounts)")
                continue
            ccy_note = (lambda v: f"{prop['currency']} {v:,.2f}") if prop["currency"] != "AUD" else (lambda v: "")
            line(f"{prop['name']} ({prop['country_name']})", sheet=prop["sheet"], bold=True,
                 note=f"{prop['weeks_rented']} weeks rented of {prop['weeks_available']} available")
            line("Gross rent", prop["rent"] * rt, sheet=prop["sheet"], note=ccy_note(prop["rent"]))
            line("Deductible expenses", -prop["expenses_total"] * rt, sheet=prop["sheet"],
                 note=ccy_note(prop["expenses_total"]))
            tax = None
            if prop["spanish_tax"]:
                tax = (prop["spanish_tax_h2"] + prop["spanish_tax_h1"]) * rt
                tot_tax += tax
            line("Net rent before depreciation", (prop["rent"] - prop["expenses_total"]) * rt, tax, prop["sheet"],
                 "Foreign tax: Spanish 24% of gross rent - see calendar split" if tax is not None
                 else ("" if prop["country"] == "AU" else f"Foreign tax: from the {prop['country_name']} returns"))
            tot_rent += prop["rent"] * rt
            tot_exp += prop["expenses_total"] * rt
        if len(props) > 1:
            line("All rental properties - net before depreciation", tot_rent - tot_exp, tot_tax or None, bold=True,
                 note=f"Rent A${tot_rent:,.2f} less expenses A${tot_exp:,.2f}"
                 + (" (excludes properties with no rate)" if missing_rate else ""))
    else:
        line(prop_src or "Not available", note="Property figures not included")

    sec("2. Investment income (foreign)")
    if income is not None and not income.empty:
        foreign = income[income["Country"] != "AU"]
        for (typ, ctry), g in foreign.groupby(["Type", "Country"]):
            line(f"{typ} - {ctry}", g["Gross A$"].sum(), g["Tax withheld A$"].sum() or None, "Investment income",
                 f"{len(g)} payment(s)")
        line("Total foreign investment income", foreign["Gross A$"].sum(), foreign["Tax withheld A$"].sum(),
             bold=True)
        au = income[income["Country"] == "AU"]
        if not au.empty:
            line("Australian (pre-filled by ATO, reference only)", au["Gross A$"].sum(), sheet="Investment income")
    else:
        line("Nothing recorded for this year")

    sec("3. Sales and maturities")
    if gains is not None and not gains.empty:
        gl = gains["Gain / (loss) A$"]
        line("Gains", gl[gl > 0].sum(), sheet="Gains and losses", note=f"{int((gl > 0).sum())} disposal(s)")
        line("Losses", gl[gl < 0].sum(), sheet="Gains and losses", note=f"{int((gl < 0).sum())} disposal(s)")
        line("Net gain / (loss)", gl.sum(), bold=True,
             note="Before any discount; held > 12 months shown per line")
    else:
        line("No sales or maturities recorded for this year")

    r += 1
    _w(ov, r, 1, "Points for the tax agent", font=BOLD)
    r += 1
    for p in points or ["None flagged."]:
        _w(ov, r, 1, "•")
        _w(ov, r, 2, p, wrap=True)
        ov.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ov.row_dimensions[r].height = 30
        r += 1
    _widths(ov, [34, 46, 14, 16, 18, 50])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), points


def rates_used_list(fy_year, props, rate_of, income, gains):
    rows = []
    for ccy in sorted({p["currency"] for p in props} - {"AUD"}):
        rt = rate_of(ccy)
        if rt:
            names = ", ".join(p["name"] for p in props if p["currency"] == ccy)
            rows.append((f"{fy_name(fy_year)} average", ccy, rt, f"Rental: {names} - ATO annual average "
                         f"({1 / rt:.4f} {ccy} per A$ = average of the Reserve Bank daily rates)"))
    if income is not None and not income.empty:
        for _, x in income[income["currency"].str.upper() != "AUD"].iterrows():
            payer = x["payer"] if isinstance(x["payer"], str) and x["payer"] else (x["portfolio"] or "")
            rows.append((x["div_date"], x["currency"], x["AUD rate"], f"{x['Type']} - {payer}"))
    if gains is not None and not gains.empty:
        for _, x in gains.iterrows():
            if str(x["Currency"]).upper() == "AUD":
                continue
            if pd.notna(x.get("Cost FX")) and x.get("Acquired") is not None:
                rows.append((x["Acquired"], x["Currency"], x["Cost FX"], f"Acquired: {x['Asset']}"))
            if pd.notna(x.get("Proceeds FX")):
                rows.append((x["Disposed"], x["Currency"], x["Proceeds FX"], f"Disposed: {x['Asset']}"))
    # one line per date+currency
    out, seen = [], {}
    for d, c, v, u in rows:
        k = (str(d), c)
        if k in seen:
            if u not in out[seen[k]][3]:
                out[seen[k]][3] += "; " + u
            continue
        seen[k] = len(out)
        out.append([d, c, v, u])
    head = [x for x in out if isinstance(x[0], str)]
    rest = sorted([x for x in out if not isinstance(x[0], str)], key=lambda x: (pd.Timestamp(x[0]), x[1]))
    return head + rest


# ─────────────────────────── page ───────────────────────────────────────────

def render_pack_page(pg, aud_rate_on, aud_avg_for_fy):
    st.header("📦 Accountant pack")
    st.caption("One Excel workbook per Australian financial year: rental properties, foreign interest, dividends "
               "and coupons, sales and maturities, and the exchange rates used.")
    cur = current_fy()
    years = list(range(cur, cur - 6, -1))
    fy_year = st.selectbox("Financial year", years, index=1,
                           format_func=lambda y: f"{fy_name(y)} (1 Jul {y - 1} - 30 Jun {y})"
                           + (" - in progress" if y == cur else ""), key="pack_fy")
    with st.spinner("Gathering figures…"):
        props, prop_src = property_for_fy(pg, fy_year)
        _rates = {}

        def rate_of(ccy):
            if ccy not in _rates:
                rt = nw_zarpia.rate_for(aud_avg_for_fy, fy_year, ccy)
                _rates[ccy] = round(rt, 6) if rt else None
            return _rates[ccy]
        income = income_for_fy(pg, fy_year, aud_rate_on)
        gains = gains_for_fy(pg, fy_year, aud_rate_on)

    c1, c2, c3 = st.columns(3)
    priced = [p for p in props if rate_of(p["currency"])]
    if priced:
        net = sum((p["rent"] - p["expenses_total"]) * rate_of(p["currency"]) for p in priced)
        c1.metric("Net rent (A$)", f"{net:,.0f}",
                  help=f"{len(priced)} propert{'y' if len(priced) == 1 else 'ies'}: "
                  + ", ".join(p["name"] for p in priced))
    else:
        c1.metric("Net rent (A$)", "-", help=prop_src)
    foreign = income[income["Country"] != "AU"] if not income.empty else income
    c2.metric("Foreign investment income (A$)", f"{foreign['Gross A$'].sum():,.0f}" if not income.empty else "0",
              help=f"{len(foreign)} payment(s)")
    c3.metric("Net gain / (loss) (A$)", f"{gains['Gain / (loss) A$'].sum():,.0f}" if not gains.empty else "0",
              help=f"{len(gains)} disposal(s)")
    if not props:
        st.caption(f"Rental properties: {prop_src}.")
    else:
        st.caption(f"Rental properties from {prop_src}: " + ", ".join(p["name"] for p in props) + ".")
    if fy_year == cur:
        st.warning("This financial year isn't over yet - the pack shows figures to date.")

    data, points = build_workbook(fy_year, props, prop_src, rate_of, income, gains,
                                  rates_used_list(fy_year, props, rate_of, income, gains))
    st.download_button(f"⬇️ Download {fy_name(fy_year)} accountant pack (Excel)", data,
                       file_name=f"Claudio_Conca_{fy_name(fy_year)}_accountant_pack.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary", key="pack_dl")
    if points:
        with st.expander(f"Points for the tax agent ({len(points)})", expanded=True):
            for p in points:
                st.markdown(f"- {p}")
    st.caption("Before sending: check the Income and Cost base & gains pages are complete for the year "
               "(e.g. Trade Republic interest, BPM bond parcels).")
